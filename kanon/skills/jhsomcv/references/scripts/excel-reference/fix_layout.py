#!/usr/bin/env python3
"""Size every block by measurement, not estimate.

Pass 1 (grow):   render, find blocks whose text prints outside the block or collides
                 with the next one, give each one more row, repeat until clean.
Pass 2 (shrink): try taking a row back off every block that was grown, all at once;
                 restore only the ones that break. This is what stops the document
                 filling up with stray blank lines.

What it needs, none of which is in this repository:

  a build command   Rebuilds the workbook from the per-block row overrides in EXTRA
                    and renders it to PDF. The renderer is specific to the CV being
                    built, so supply your own with --build or JHSOMCV_BUILD, e.g.
                        --build "python3 render.py && bash mkpreview.sh"
                    It runs through the shell in the current directory and must read
                    EXTRA, write XLSX and produce PDF.
  openpyxl          to read the workbook.
  pdftotext         from poppler, to read the rendered PDF.

Usage:
    python3 fix_layout.py --build CMD [--xlsx cv.xlsx] [--pdf preview.pdf]
                          [--extra extra_rows.json] [--reset]
"""
import argparse, json, os, re, shutil, subprocess, sys

def die(msg):
    sys.exit('fix_layout.py: ' + msg)

try:
    import openpyxl
except ImportError:
    die('openpyxl is not installed (pip install openpyxl)')
from audit_layout import check_overlaps
from check_clipping import check as geom_check

XLSX = os.environ.get('JHSOMCV_XLSX', 'cv.xlsx')
PDF = os.environ.get('JHSOMCV_PDF', 'preview.pdf')
EXTRA = os.environ.get('JHSOMCV_EXTRA', 'extra_rows.json')
BUILD = os.environ.get('JHSOMCV_BUILD')

def key(t): return re.sub(r'\s+', ' ', str(t))   # full text: a 90-char prefix collides between duplicate-prefix entries

def build():
    """Run the caller's build command and confirm it produced the workbook and the PDF."""
    if not BUILD:
        die('no build command. Pass --build CMD or set JHSOMCV_BUILD; the command must '
            'read %s, write %s and render %s. The renderer is CV-specific and is not '
            'part of this repository.' % (EXTRA, XLSX, PDF))
    r = subprocess.run(BUILD, shell=True, capture_output=True, text=True)
    if r.returncode:
        die('build command failed (exit %d): %s\n%s' % (r.returncode, BUILD, r.stderr.strip()[-2000:]))
    for f, what in ((XLSX, 'workbook'), (PDF, 'PDF')):
        if not os.path.exists(f):
            die('build command ran but did not produce the %s %s: %s' % (what, f, BUILD))

def load():
    if not os.path.exists(EXTRA):
        return {}
    with open(EXTRA) as f:
        return json.load(f)

def save(e):
    with open(EXTRA, 'w') as f:
        json.dump(e, f, indent=1)

def pdf_lines(pdf):
    out = subprocess.run(['pdftotext', '-bbox', pdf, '-'], capture_output=True, text=True).stdout
    pages = []
    for page in out.split('<page ')[1:]:
        w = [(float(a), float(b), t) for a, b, c, d, t in
             re.findall(r'<word xMin="([\d.]+)" yMin="([\d.]+)" xMax="([\d.]+)" yMax="([\d.]+)">([^<]*)</word>', page)]
        pts = sorted((y, x, t) for x, y, t in w if 20 <= y <= 745 and set(t) - set('_'))
        ys, L = [], {}
        for y, x, t in pts:
            if ys and y - ys[-1] <= 2.0: L[ys[-1]].append((x, t))
            else: ys.append(y); L[y] = [(x, t)]
        pages.append([(y, ' '.join(t for _, t in sorted(L[y]))) for y in ys])
    return pages

def offenders():
    """Keys of blocks that need another row."""
    bad = set()
    ws = openpyxl.load_workbook(XLSX).active
    cells = [c.value for row in ws.iter_rows() for c in row
             if isinstance(c.value, str) and len(c.value) > 20]
    for coord, rows, over, txt in geom_check(XLSX, PDF):
        v = ws[coord].value
        if isinstance(v, str): bad.add(key(v))
    # Measurement beats estimate. The geometric check reads the rendered PDF and is
    # ground truth for any block it can trace; the analytic wrap model is an estimate
    # that runs about one line long on paragraphs of eight lines or more. Union the
    # two and every long narrative block keeps a row it does not need, which is
    # exactly the "random extra space under the paragraph" complaint. So the analytic
    # result is consulted only for blocks the geometry could NOT trace.
    _, _unver = geom_check(XLSX, PDF, report_unverified=True)
    _unkeys = {key(ws[c].value) for c, _t in _unver if isinstance(ws[c].value, str)}
    if _unkeys:
        from audit_layout import audit as _analytic
        for kind, coord, _msg in _analytic(XLSX)[0]:
            if kind == 'CLIPPED':
                v = ws[coord].value
                if isinstance(v, str) and key(v) in _unkeys: bad.add(key(v))
    ovl = [o for o in check_overlaps(PDF) if o[0] == 'COLLISION']
    if ovl:
        pages = pdf_lines(PDF)
        for sev, pno, y, g, modal, a, b in ovl:
            lines = pages[pno - 1]
            idx = next((i for i, (yy, tt) in enumerate(lines) if abs(yy - y) < 0.5), None)
            ctx = ' '.join(tt for _, tt in lines[max(0, idx - 1):idx + 1]) if idx is not None else a
            frag = re.sub(r'\s+', ' ', ctx).replace('&quot;', '"').replace('&amp;', '&').strip()
            for tail in (frag[-28:], re.sub(r'\s+', ' ', a).strip()[-18:]):
                hit = [v for v in cells if re.sub(r'\s+', ' ', v).rstrip().endswith(tail)]
                if hit:
                    bad.add(key(min(hit, key=len))); break
    return bad

def grow():
    for i in range(8):
        build()
        bad = offenders()
        print('  grow pass %d: %d block(s) need more room' % (i, len(bad)))
        if not bad: return
        e = load()
        for k in bad: e[k] = e.get(k, 0) + 1
        save(e)

def trim():
    """Take back rows the model over-allotted. Detected geometrically: a block whose
    text ends more than a full row above its bottom edge has a spare row. Trim all of
    them, then restore only the ones that break. Iterate: each pass restores the
    blocks that broke, which frees the geometry to reveal slack the pass before it
    was masking, so a low pass cap silently leaves spare rows in the document."""
    for i in range(30):
        _, slack = geom_check(XLSX, PDF, report_slack=True)
        if not slack:
            print('  trim pass %d: no spare rows' % i); return
        ws = openpyxl.load_workbook(XLSX).active
        e = load()
        touched = {}
        for coord, rows, spare, txt in slack:
            v = ws[coord].value
            if isinstance(v, str):
                k = key(v); touched[k] = spare; e[k] = e.get(k, 0) - spare
        save(e); build()
        bad = offenders()
        if bad:
            for k in bad:
                if k in touched:
                    e[k] = e.get(k, 0) + touched[k]
            save(e); build()
            if offenders():                  # still broken: undo the whole pass
                for k, sp in touched.items():
                    e[k] = e.get(k, 0) + sp
                save(e); build()
                print('  trim pass %d: reverted' % i); return
        removed = len(touched) - len(bad & set(touched))
        print('  trim pass %d: removed %d spare row(s)' % (i, removed))
        if removed == 0:
            return

def shrink():
    for i in range(8):
        e = load()
        cand = [k for k, v in e.items() if v > 0]
        if not cand: return
        trial = {k: (v - 1 if v > 0 else v) for k, v in e.items()}
        save(trial)
        build()
        bad = offenders()
        if not bad:
            print('  shrink pass %d: removed a row from %d block(s), still clean' % (i, len(cand)))
            continue
        # restore only the blocks that actually needed the row, and stop
        for k in bad:
            trial[k] = trial.get(k, 0) + 1
        save({k: v for k, v in trial.items() if v > 0})
        build()
        if not offenders():
            print('  shrink pass %d: settled, %d block(s) keep an extra row' % (i, len(bad)))
            return
        # restoring was not enough: go back to the last known-good state
        save(e); build()
        print('  shrink pass %d: reverted' % i)
        return

def main():
    global XLSX, PDF, EXTRA, BUILD
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('--build', default=BUILD, help='shell command that rebuilds XLSX from EXTRA and renders PDF (or set JHSOMCV_BUILD)')
    ap.add_argument('--xlsx', default=XLSX, help='workbook the build command writes (default %(default)s)')
    ap.add_argument('--pdf', default=PDF, help='PDF the build command renders (default %(default)s)')
    ap.add_argument('--extra', default=EXTRA, help='per-block row overrides, read by the build command (default %(default)s)')
    ap.add_argument('--reset', action='store_true', help='discard the existing row overrides first')
    a = ap.parse_args()
    XLSX, PDF, EXTRA, BUILD = a.xlsx, a.pdf, a.extra, a.build
    if not BUILD:
        ap.error('no build command: pass --build CMD or set JHSOMCV_BUILD (see --help)')
    if shutil.which('pdftotext') is None:
        die('pdftotext is not on PATH; install poppler')
    if a.reset and os.path.exists(EXTRA):
        os.remove(EXTRA)
    grow()
    shrink()
    trim()
    build()
    print('final:', len(offenders()), 'offender(s);', sum(load().values()), 'extra row(s) across', len(load()), 'block(s)')

if __name__ == '__main__':
    main()